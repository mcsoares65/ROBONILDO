"""Classificação oficial dos cartuchos executados pelo console motor.py.

Critério único de ordenação para entrada, saída e cruzado:

  acumulado = resultado - abs(drawdown)

Como o drawdown é armazenado com sinal negativo, a mesma fórmula também
pode ser escrita como ``resultado + drawdown``. Quanto maior o lucro e menor
o drawdown, maior o acumulado. Não há filtro de elegibilidade, pontos por dia
nem portão estatístico: a posição é só o acumulado. As métricas multitemporais
continuam sendo calculadas para auditoria, mas não definem a posição.

Modos: E (entradas × saída titular), S (saídas × entrada titular),
C (todas as combinações) e A (Análise: roda o cruzado completo e grava uma
planilha .xlsx com rankings e apuração dia/mês/ano por estratégia) e
D (Diagnóstico de cenários: R$ e operações de cada entrada e saída em cada
cenário, só leitura; incorpora o antigo diagnostico_cenarios.py).
"""

from __future__ import annotations

import csv
import importlib.util
import io
import math
import re
import subprocess
import contextlib
from collections import defaultdict
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import median
from typing import Callable, Iterable, Optional

import cenario as cen_mod
import configuracao as cfg
from motor import Candle, MotorRobonildo, construir_row

# Versionamento próprio deste programa. Alterações em regras de validação,
# ordenação ou apresentação da classificação incrementam esta constante sem
# alterar a versão operacional do Robonildo definida em configuracao.py.
VERSAO_CLASSIFICACAO = "C001"

_ANSI_RE = re.compile(r"\x1b\[[0-9;]*m")


def _git_sha_atual() -> str:
    """SHA do commit atual do repositório, ou "sem-git" se não estiver
    rodando dentro de um clone git (ex: pasta copiada manualmente). Usado só
    para carimbar `logs/classificacao_historico.md` - nunca falha o
    programa, silenciosamente vira "sem-git"."""
    try:
        return (
            subprocess.check_output(
                ["git", "rev-parse", "HEAD"],
                cwd=Path(__file__).resolve().parent,
                stderr=subprocess.DEVNULL,
            )
            .decode()
            .strip()
        )
    except Exception:
        return "sem-git"


def _registrar_historico(titulo: str, texto: str) -> None:
    """Acrescenta (NUNCA sobrescreve) o resultado desta rodada em
    `logs/classificacao_historico.md`, carimbado com data/hora e o SHA do
    commit atual do repositório.

    Existe para que nenhuma promoção de titular dependa de alguém (humano
    ou IA) colar manualmente um resultado no chat ou num PR - o próprio
    `classificacao.py` já deixa a prova gravada, atrelada ao commit exato
    em que rodou (compliance.md, Regra 4: nenhum resultado vale sem ter
    sido rodado de novo contra o motor atual - o SHA aqui é o que permite
    conferir isso depois, sem confiar na palavra de quem rodou).
    """
    pasta_logs = Path(cfg.PASTA_LOGS_AUDITORIA)
    pasta_logs.mkdir(parents=True, exist_ok=True)
    caminho = pasta_logs / "classificacao_historico.md"
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    sha = _git_sha_atual()
    texto_limpo = _ANSI_RE.sub("", texto).rstrip()
    bloco = f"\n## {agora} — commit `{sha}`\n\n{titulo}\n\n```\n{texto_limpo}\n```\n"
    with caminho.open("a", encoding="utf-8") as f:
        f.write(bloco)


@contextlib.contextmanager
def _capturar_e_imprimir():
    """Deixa o bloco `with` imprimir normalmente no terminal E devolve o
    texto capturado (sem códigos ANSI removidos ainda - isso é feito em
    `_registrar_historico`) para gravar no histórico."""
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        yield buf
    texto = buf.getvalue()
    print(texto, end="")

try:
    import colorama
    colorama.init()
    VERDE = colorama.Fore.LIGHTGREEN_EX
    VERMELHO = colorama.Fore.RED
    RESET = colorama.Style.RESET_ALL
except ImportError:
    VERDE = VERMELHO = RESET = ""
# Azul celeste para a linha da estrategia titular - fora da paleta padrao do
# colorama (so tem os 8 nomes classicos), por isso usa o codigo ANSI 256
# direto (117 = azul celeste claro). Mesma tecnica ja usada no marcador de
# posicao em principal.py.
AZUL_CELESTE = "\x1b[38;5;117m" if RESET else ""


def _progresso(etapa: str, atual: int, total: int, detalhe: str = "") -> None:
    """Atualiza uma única linha do terminal sem poluir a classificação."""
    percentual = 100.0 * atual / total if total else 100.0
    complemento = f" | {detalhe}" if detalhe else ""
    linha = (
        f"[{etapa}] {atual:>{len(str(max(total, 1)))}}/{total} "
        f"({percentual:6.2f}%){complemento}"
    )
    print("\r" + linha.ljust(110), end="", flush=True)   # ljust apaga resto da linha anterior
    if atual >= total:
        print()


def carregar_csv(caminho: Path) -> list[Candle]:
    candles = []
    tamanho_total = max(caminho.stat().st_size, 1)
    bytes_lidos = 0
    linhas_lidas = 0
    print(f"[LEITURA] Abrindo {caminho.resolve()}")
    with caminho.open("rb") as arquivo:
        for linha_bruta in arquivo:
            bytes_lidos += len(linha_bruta)
            linhas_lidas += 1
            linha = next(csv.reader([linha_bruta.decode("latin1")], delimiter=";"))
            if len(linha) < 7:
                pass
            else:
                try:
                    horario = datetime.strptime(
                        f"{linha[1]} {linha[2]}", "%d/%m/%Y %H:%M:%S"
                    )
                    valores = [float(v.replace(".", "").replace(",", ".")) for v in linha[3:7]]
                except ValueError:
                    pass
                else:
                    quantidade = None
                    if len(linha) > 8 and linha[8].strip():
                        try:
                            quantidade = float(linha[8].replace(".", "").replace(",", "."))
                            # Histórico persistido legado usava zero como
                            # placeholder, não como quantidade observada.
                            if not math.isfinite(quantidade) or quantidade <= 0:
                                quantidade = None
                        except ValueError:
                            pass
                    candles.append(Candle(horario, *valores, quantidade))
            if linhas_lidas % 500 == 0:
                _progresso("LEITURA", min(bytes_lidos, tamanho_total), tamanho_total,
                           f"{len(candles)} candles válidos")

    _progresso("LEITURA", tamanho_total, tamanho_total, f"{len(candles)} candles válidos")

    por_horario = {c.horario: c for c in candles}
    ordenados = sorted(por_horario.values(), key=lambda c: c.horario)
    if not ordenados:
        raise ValueError("O histórico não contém candles válidos.")
    return ordenados


def preparar_aquecimento(candles_avaliacao: list[Candle]) -> tuple[list[Candle], int]:
    """Completa históricos curtos com candles anteriores, sem pontuá-los."""
    if len(candles_avaliacao) >= 66:
        return candles_avaliacao, 0

    caminho = Path(cfg.LAB_CAMINHO_HISTORICO_AQUECIMENTO)
    if not caminho.exists():
        raise ValueError(
            "O arquivo avaliado não possui 66 candles para calcular os indicadores e "
            f"o histórico de aquecimento não foi encontrado em '{caminho}'."
        )
    print(f"[AQUECIMENTO] Histórico curto detectado; buscando candles anteriores em {caminho}.")
    historico = carregar_csv(caminho)
    inicio = candles_avaliacao[0].horario
    anteriores = [candle for candle in historico if candle.horario < inicio][-65:]
    if len(anteriores) < 65:
        raise ValueError(
            "Histórico curto: são necessários 65 candles anteriores para aquecer "
            f"os indicadores, mas somente {len(anteriores)} foram encontrados."
        )
    por_horario = {c.horario: c for c in anteriores + candles_avaliacao}
    combinados = sorted(por_horario.values(), key=lambda c: c.horario)
    print(f"[AQUECIMENTO] {len(anteriores)} candles anteriores adicionados; eles não serão pontuados.")
    return combinados, len(anteriores)


def _importar_arquivo(caminho: Path, indice):
    nome_modulo = f"robonildo_lab_{indice}_{abs(hash(caminho.resolve()))}"
    spec = importlib.util.spec_from_file_location(nome_modulo, caminho)
    if spec is None or spec.loader is None:
        raise ImportError(f"Não foi possível importar {caminho}")
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo



@dataclass
class CartuchoEntrada:
    nome: str
    caminho: Path
    gerar_sinal: Callable
    titular: bool = False


@dataclass
class CartuchoSaida:
    nome: str
    caminho: Path
    avaliar_saida: Callable
    titular: bool = False


# Compat: nome antigo usado em trechos legados
EstrategiaCarregada = CartuchoEntrada


def _listar_py(pasta: Path) -> list[Path]:
    """Arquivos .py de cartucho diretamente numa pasta (não desce em subpastas).

    Ignora nomes que começam com `_` (reservados, Regra 7) e arquivos de teste
    `test_*.py`, que ficam ao lado das candidatas no laboratório. Auxiliares de
    cartucho moram em `auxiliar/` (subpasta, nunca varrida)."""
    if not pasta.exists():
        return []
    return sorted(
        p for p in pasta.glob("*.py")
        if not p.stem.startswith(("_", "test_"))
    )


def listar_cartuchos_disco() -> dict:
    """
    Estrutura oficial (sem pasta laboratorio):

      estrategia/entrada/titular/   → exatamente 1 .py  (entrada de produção)
      estrategia/entrada/*.py      → candidatas de ranking de ENTRADA

      estrategia/saida/titular/     → exatamente 1 .py  (saída de produção)
      estrategia/saida/*.py        → candidatas de ranking de SAÍDA

    Rankings:
      1) entradas × saída titular
      2) saída candidatas × entrada titular
      3) cruzado (todas entradas × todas saídas) — ranking principal
    """
    raiz = Path(__file__).resolve().parent
    pasta_entrada = raiz / "estrategia" / "entrada"
    pasta_saida = raiz / "estrategia" / "saida"
    pasta_entrada_titular = pasta_entrada / "titular"
    pasta_saida_titular = pasta_saida / "titular"

    return {
        "raiz": raiz,
        "pasta_entrada": pasta_entrada,
        "pasta_saida": pasta_saida,
        "pasta_entrada_titular": pasta_entrada_titular,
        "pasta_saida_titular": pasta_saida_titular,
        "entrada_titular": _listar_py(pasta_entrada_titular),
        "entrada_candidatas": _listar_py(pasta_entrada),
        "saida_titular": _listar_py(pasta_saida_titular),
        "saida_candidatas": _listar_py(pasta_saida),
    }


def descobrir_entradas() -> tuple[list[CartuchoEntrada], list[tuple[str, str]]]:
    """Todas as entradas: titular + candidatas na raiz de entrada/."""
    inv = listar_cartuchos_disco()
    titulares = inv["entrada_titular"]
    if len(titulares) != 1:
        raise RuntimeError(
            "A pasta estrategia/entrada/titular/ deve conter exatamente um arquivo. "
            f"Encontrados: {[p.name for p in titulares]}"
        )
    caminho_titular = titulares[0]

    caminhos = [caminho_titular]
    vistos = {caminho_titular.stem.casefold()}
    for caminho in inv["entrada_candidatas"]:
        if caminho.stem.casefold() in vistos:
            continue
        vistos.add(caminho.stem.casefold())
        caminhos.append(caminho)

    carregadas = []
    falhas = []
    for indice, caminho in enumerate(caminhos):
        try:
            modulo = _importar_arquivo(caminho, f"ent_{indice}")
        except Exception as erro:
            if caminho.resolve() == caminho_titular.resolve():
                raise
            falhas.append(
                (caminho.stem, f"falha na importação: {type(erro).__name__}: {erro}")
            )
            continue
        gerar = getattr(modulo, "gerar_sinal", None)
        if callable(gerar):
            carregadas.append(CartuchoEntrada(
                nome=caminho.stem,
                caminho=caminho,
                gerar_sinal=gerar,
                titular=(caminho.resolve() == caminho_titular.resolve()),
            ))
        else:
            falhas.append((caminho.stem, "função gerar_sinal(row) ausente"))
    return carregadas, falhas


def descobrir_saidas() -> tuple[list[CartuchoSaida], list[tuple[str, str]]]:
    """Todas as saídas: titular + candidatas na raiz de saida/."""
    inv = listar_cartuchos_disco()
    titulares = inv["saida_titular"]
    if len(titulares) != 1:
        raise RuntimeError(
            "A pasta estrategia/saida/titular/ deve conter exatamente um arquivo. "
            f"Encontrados: {[p.name for p in titulares]}"
        )
    caminho_titular = titulares[0]

    caminhos = [caminho_titular]
    vistos = {caminho_titular.stem.casefold()}
    for caminho in inv["saida_candidatas"]:
        if caminho.stem.casefold() in vistos:
            continue
        vistos.add(caminho.stem.casefold())
        caminhos.append(caminho)

    carregadas = []
    falhas = []
    for indice, caminho in enumerate(caminhos):
        try:
            modulo = _importar_arquivo(caminho, f"sai_{indice}")
        except Exception as erro:
            if caminho.resolve() == caminho_titular.resolve():
                raise
            falhas.append(
                (caminho.stem, f"falha na importação: {type(erro).__name__}: {erro}")
            )
            continue
        avaliar = getattr(modulo, "avaliar_saida", None)
        if callable(avaliar):
            carregadas.append(CartuchoSaida(
                nome=caminho.stem,
                caminho=caminho,
                avaliar_saida=avaliar,
                titular=(caminho.resolve() == caminho_titular.resolve()),
            ))
        else:
            falhas.append((caminho.stem, "função avaliar_saida(row, posicao) ausente"))
    return carregadas, falhas


def descobrir_estrategias() -> tuple[list[CartuchoEntrada], list[tuple[str, str]]]:
    """Compat: retorna só entradas (titular + candidatas)."""
    return descobrir_entradas()


def descobrir_cartucho_saida():
    """Compat: retorna (nome, callable) da saída titular."""
    saidas, _ = descobrir_saidas()
    titular = next((s for s in saidas if s.titular), None)
    if titular is None:
        raise RuntimeError("Nenhuma saída titular encontrada em estrategia/saida/titular/.")
    return titular.nome, titular.avaliar_saida




def preparar_rows(candles: list[Candle], dias_avaliacao: Optional[set] = None) -> list[Optional[dict]]:
    """Indicadores por candle. Com dias_avaliacao, só calcula os candles desses
    dias (os anteriores continuam servindo de aquecimento, sem serem pontuados)."""
    rows = [None] * len(candles)
    alvo = [i for i in range(65, len(candles))
            if dias_avaliacao is None or candles[i].horario.date() in dias_avaliacao]
    total = len(alvo)
    for concluido, indice in enumerate(alvo, 1):
        rows[indice] = construir_row(candles[:indice + 1])
        if concluido % 250 == 0 or concluido == total:
            _progresso("INDICADORES", concluido, total)
    return rows


_RE_DATA = re.compile(r"(\d{2})/(\d{2})/(\d{4})|(\d{4})-(\d{2})-(\d{2})")


def _datas_do_texto(texto: str) -> list:
    achadas = []
    for m in _RE_DATA.finditer(texto):
        if m.group(1):
            achadas.append(date(int(m.group(3)), int(m.group(2)), int(m.group(1))))
        else:
            achadas.append(date(int(m.group(4)), int(m.group(5)), int(m.group(6))))
    return achadas


def interpretar_periodo(texto: Optional[str], dias: list) -> tuple[list, str]:
    """Converte o texto do período em (dias selecionados, rótulo).

    Aceita (vazio ou 'tudo' = todo o histórico):
      01/03/2026 31/03/2026   intervalo de datas (também aaaa-mm-dd, ou com ':' / 'a' / 'ate')
      01/03/2026              deste dia até o fim do histórico
      ultimos 40              os 40 últimos pregões
      mes 2026-03             um mês (também 03/2026)
    Levanta ValueError com mensagem clara se o texto não for entendido ou o
    período não tiver nenhum pregão.
    """
    bruto = (texto or "").strip()
    t = bruto.lower().replace("é", "e").replace("ú", "u").replace("ê", "e")
    if t in ("", "tudo", "todo", "todos"):
        return list(dias), "tudo"

    m = re.fullmatch(r"(?:ultimos|ultimo|u)\s*(\d+)", t)
    if m:
        n = int(m.group(1))
        if n < 1:
            raise ValueError("'ultimos N' precisa de N >= 1")
        sel = list(dias)[-n:]
        return sel, f"ultimos {len(sel)} pregoes"

    m = re.fullmatch(r"(?:mes\s*)?(\d{4})-(\d{2})", t) or re.fullmatch(r"(?:mes\s*)?(\d{2})/(\d{4})", t)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        ano, mes = (a, b) if a > 31 else (b, a)
        if not 1 <= mes <= 12:
            raise ValueError(f"mês inválido em '{bruto}'")
        sel = [d for d in dias if d.year == ano and d.month == mes]
        if not sel:
            raise ValueError(f"nenhum pregão em {mes:02d}/{ano} no histórico "
                             f"({dias[0].strftime('%d/%m/%Y')} a {dias[-1].strftime('%d/%m/%Y')})")
        return sel, f"{ano}-{mes:02d}"

    try:
        datas = _datas_do_texto(t)
    except ValueError as erro:
        raise ValueError(f"data inválida em '{bruto}': {erro}") from erro
    if len(datas) in (1, 2):
        ini = datas[0]
        fim = datas[1] if len(datas) == 2 else dias[-1]
        if fim < ini:
            ini, fim = fim, ini
        sel = [d for d in dias if ini <= d <= fim]
        if not sel:
            raise ValueError(f"nenhum pregão entre {ini.strftime('%d/%m/%Y')} e {fim.strftime('%d/%m/%Y')} "
                             f"(histórico: {dias[0].strftime('%d/%m/%Y')} a {dias[-1].strftime('%d/%m/%Y')})")
        return sel, f"{sel[0].isoformat()}_{sel[-1].isoformat()}"
    raise ValueError(
        f"não entendi o período '{bruto}'. Exemplos: 01/03/2026 31/03/2026 | ultimos 40 | mes 2026-03 | Enter = tudo")


def validar_contrato_saida(cartucho: CartuchoSaida, rows: list[Optional[dict]]) -> tuple[bool, str]:
    """Valida o contrato V445 antes de permitir que a saída participe do ranking.

    A V445 não fornece proteção implícita. Portanto, uma saída compatível deve
    retornar um ``dict`` na abertura e fornecer um stop inicial finito e
    coerente tanto para COMPRA quanto para VENDA. O alvo pode ser ``None``
    quando o próprio cartucho optar explicitamente por trabalhar sem alvo.
    """
    row = next((item for item in reversed(rows) if item is not None), None)
    if row is None:
        return False, "não há indicadores suficientes para validar a abertura"

    ohlc = row.get("ohlc_recentes") or ()
    janela = ohlc[-5:]
    if len(janela) < 5:
        return False, "row de validação não contém cinco candles em ohlc_recentes"

    minima = min(float(c["Minimo"]) for c in janela)
    maxima = max(float(c["Maximo"]) for c in janela)
    if not math.isfinite(minima) or not math.isfinite(maxima) or maxima <= minima:
        return False, "janela OHLC inválida para o teste de abertura"

    # O ponto médio evita reprovar artificialmente uma fórmula estrutural
    # quando o último fechamento coincide exatamente com um extremo da janela.
    entrada = (minima + maxima) / 2.0
    for lado in ("COMPRA", "VENDA"):
        posicao = {
            "lado": lado,
            "entrada": entrada,
            "candles_decorridos": 0,
            "maxima_desde_entrada": entrada,
            "minima_desde_entrada": entrada,
            "resultado_flutuante_pts": 0.0,
        }
        try:
            resposta = cartucho.avaliar_saida(row, posicao)
        except Exception as erro:
            return False, f"falhou na abertura de {lado}: {type(erro).__name__}: {erro}"

        if not isinstance(resposta, dict):
            return False, (
                f"retornou {type(resposta).__name__} na abertura de {lado}; "
                "a V445 exige dict com novo_stop e novo_alvo"
            )
        if bool(resposta.get("fechar", False)):
            return False, f"solicitou fechamento durante a abertura de {lado}"
        if "novo_stop" not in resposta or "novo_alvo" not in resposta:
            return False, f"retorno de {lado} não contém novo_stop e novo_alvo"

        try:
            stop = float(resposta["novo_stop"])
        except (TypeError, ValueError):
            return False, f"novo_stop ausente ou não numérico na abertura de {lado}"
        if not math.isfinite(stop):
            return False, f"novo_stop não finito na abertura de {lado}"
        if lado == "COMPRA" and stop >= entrada:
            return False, "stop inicial de COMPRA não está abaixo da entrada"
        if lado == "VENDA" and stop <= entrada:
            return False, "stop inicial de VENDA não está acima da entrada"

        alvo_bruto = resposta.get("novo_alvo")
        if alvo_bruto is not None:
            try:
                alvo = float(alvo_bruto)
            except (TypeError, ValueError):
                return False, f"novo_alvo não numérico na abertura de {lado}"
            if not math.isfinite(alvo):
                return False, f"novo_alvo não finito na abertura de {lado}"
            if lado == "COMPRA" and alvo <= entrada:
                return False, "alvo inicial de COMPRA não está acima da entrada"
            if lado == "VENDA" and alvo >= entrada:
                return False, "alvo inicial de VENDA não está abaixo da entrada"

    return True, ""


def _ultimo_candle_do_dia(candles: list[Candle], indice: int) -> bool:
    return (
        indice == len(candles) - 1
        or candles[indice + 1].horario.date() != candles[indice].horario.date()
    )


def executar_jogo(
    candles: list[Candle],
    rows: list[Optional[dict]],
    estrategia,  # CartuchoEntrada
    dias_avaliacao: set,
    avaliar_saida=None,
    cenario: Optional[str] = None,
) -> list[dict]:
    """Executa o histórico diretamente no mesmo motor usado por principal.py.

    cenario: se informado (nome de cenario.py), a entrada só pode disparar em
    candles desse cenário - é o jogador escalado só para aquela situação. O
    resto (motor, saída, limites) é idêntico.

    avaliar_saida: cartucho de saída titular (estrategia/saida/) - pareado
    com TODA estratégia de entrada candidata, para que a classificação
    reflita o que rodaria em produção com a saída atual."""
    if len(candles) <= 65:
        return []

    gerar_sinal = estrategia.gerar_sinal
    if cenario:
        sinal_livre = gerar_sinal

        def gerar_sinal(row, _livre=sinal_livre, _alvo=cenario):
            return _livre(row) if cen_mod.classificar(row) == _alvo else 0

    motor = MotorRobonildo(
        gerar_sinal=gerar_sinal,
        arquivo_estado=None,
        horario_mercado_inicial=candles[65].horario,
        avaliar_saida=avaliar_saida,
    )
    for indice in range(65, len(candles)):
        candle = candles[indice]
        if candle.horario.date() not in dias_avaliacao:
            continue
        row = rows[indice]
        if row is None:
            continue
        motor.processar_candle_historico(
            candles[:indice + 1],
            ultimo_candle_do_dia=_ultimo_candle_do_dia(candles, indice),
            row=row,
        )
    return motor.operacoes_fechadas


def _profit_factor(trades: Iterable[dict]) -> float:
    valores = [t["resultado_reais"] for t in trades]
    ganhos = sum(v for v in valores if v > 0)
    perdas = abs(sum(v for v in valores if v <= 0))
    if not valores or perdas == 0:
        return float("nan")
    return ganhos / perdas


def _resumo(trades: list[dict], dias_periodo: int) -> dict:
    resultado = sum(t["resultado_reais"] for t in trades)
    dias_operados = len({t["horario_rotulo"].date() for t in trades})
    acumulado = pico = drawdown = 0.0
    for trade in trades:
        acumulado += trade["resultado_reais"]
        pico = max(pico, acumulado)
        drawdown = min(drawdown, acumulado - pico)
    return {
        "resultado": resultado,
        "diaria": resultado / dias_periodo if dias_periodo else 0.0,
        "dias_operados": dias_operados,
        "aproveitamento": resultado / dias_operados if dias_operados else 0.0,
        "drawdown": drawdown,
        "capital_minimo": cfg.MARGEM_WIN_LABORATORIO + abs(drawdown),
        "pf": _profit_factor(trades),
        "operacoes": len(trades),
        "ambiguidades": sum(1 for t in trades if t["intrabar_ambiguo"]),
    }


def _percentil_25(valores: list[float]) -> float:
    if not valores:
        return float("nan")
    valores = sorted(valores)
    posicao = 0.25 * (len(valores) - 1)
    inferior = math.floor(posicao)
    superior = math.ceil(posicao)
    if inferior == superior:
        return valores[inferior]
    fracao = posicao - inferior
    return valores[inferior] * (1 - fracao) + valores[superior] * fracao


def _pontuacao(resultado: float, capital: float, pf: float, dias: int) -> float:
    if capital <= 0 or dias <= 0 or math.isnan(pf):
        return float("nan")
    return (resultado / capital) * pf * math.sqrt(dias)


def avaliar_robustez(trades: list[dict], dias: list, tamanho: int, passo: int) -> dict:
    pontuacoes = []
    resultados = []
    for inicio in range(0, max(0, len(dias) - tamanho + 1), passo):
        janela = set(dias[inicio:inicio + tamanho])
        locais = [t for t in trades if t["horario_rotulo"].date() in janela]
        resumo = _resumo(locais, tamanho)
        score = _pontuacao(
            resumo["resultado"], resumo["capital_minimo"], resumo["pf"],
            resumo["dias_operados"],
        )
        if not math.isnan(score):
            pontuacoes.append(score)
        resultados.append(resumo["resultado"])
    return {
        "janelas": len(resultados),
        "janelas_positivas_pct": 100 * sum(v > 0 for v in resultados) / len(resultados) if resultados else 0.0,
        "pior_janela": min(resultados) if resultados else 0.0,
        "mediana_janela": median(resultados) if resultados else 0.0,
        "pontuacao_robusta": _percentil_25(pontuacoes),
    }


def _media(valores: list[float]) -> float:
    return sum(valores) / len(valores) if valores else float("nan")


def _score_periodo(trades: list[dict], quantidade_dias: int) -> float:
    """Nota simples de retorno por capital, ajustada pela recorrência."""
    if quantidade_dias <= 0:
        return float("nan")
    resumo = _resumo(trades, quantidade_dias)
    if resumo["capital_minimo"] <= 0:
        return float("nan")
    recorrencia = math.sqrt(max(resumo["dias_operados"], 1))
    return (resumo["resultado"] / resumo["capital_minimo"]) * recorrencia


def avaliar_horizontes(trades: list[dict], dias: list) -> dict:
    """Calcula as notas brutas mensal e diária, incluindo pregões sem trade."""
    por_dia = {dia: [] for dia in dias}
    for trade in trades:
        por_dia.setdefault(trade["horario_rotulo"].date(), []).append(trade)
    scores_diarios = [_score_periodo(por_dia[dia], 1) for dia in dias]

    meses = sorted({(dia.year, dia.month) for dia in dias})
    scores_mensais = []
    for ano, mes in meses:
        dias_mes = [dia for dia in dias if (dia.year, dia.month) == (ano, mes)]
        trades_mes = [
            trade for trade in trades
            if (trade["horario_rotulo"].year, trade["horario_rotulo"].month) == (ano, mes)
        ]
        scores_mensais.append(_score_periodo(trades_mes, len(dias_mes)))

    mensal = (
        0.40 * _media(scores_mensais)
        + 0.30 * median(scores_mensais)
        + 0.30 * _percentil_25(scores_mensais)
    ) if scores_mensais else float("nan")
    diario = (
        0.30 * _media(scores_diarios)
        + 0.30 * median(scores_diarios)
        + 0.40 * _percentil_25(scores_diarios)
    ) if scores_diarios else float("nan")
    return {
        "score_mensal_raw": mensal,
        "score_diario_raw": diario,
        "meses_avaliados": len(meses),
        "pregoes_avaliados": len(dias),
    }


def aplicar_pontuacao_multitemporal(resultados: list[dict], dias: list) -> dict:
    """Score ABSOLUTO individual — sem percentil entre concorrentes.

    multi = (peso_longo  × score_longo_raw
           + peso_mensal × score_mensal_raw
           + peso_diario × score_diario_raw) / pesos_válidos

    Pesos padrão (cfg): 50% longo + 30% mensal + 20% diário.
    Se um horizonte não estiver disponível, redistribui entre os ativos.

    Um colega que rode uma única estratégia isolada obtém o MESMO multi
    que obteria aqui no arsenal completo — comparável por valor, não por rank.
    """
    meses = len({(dia.year, dia.month) for dia in dias})
    usar_longo = meses >= cfg.LAB_MIN_MESES_LONGO
    usar_mensal = meses >= cfg.LAB_MIN_MESES_MENSAL and len(dias) >= cfg.LAB_MIN_PREGOES_MENSAL
    usar_diario = bool(dias)

    # Notas = scores brutos individuais (não percentil 0-100).
    for resultado in resultados:
        resultado["nota_longo"] = (
            resultado["score_longo_raw"] if usar_longo else float("nan")
        )
        resultado["nota_mensal"] = (
            resultado["score_mensal_raw"] if usar_mensal else float("nan")
        )
        resultado["nota_diario"] = (
            resultado["score_diario_raw"] if usar_diario else float("nan")
        )

    horizontes = [
        ("nota_longo", cfg.LAB_PESO_LONGO, usar_longo),
        ("nota_mensal", cfg.LAB_PESO_MENSAL, usar_mensal),
        ("nota_diario", cfg.LAB_PESO_DIARIO, usar_diario),
    ]
    peso_disponivel = sum(peso for _, peso, ativo in horizontes if ativo)
    for resultado in resultados:
        parcelas = [
            resultado[chave] * peso
            for chave, peso, ativo in horizontes
            if ativo and not math.isnan(resultado[chave])
        ]
        pesos_validos = sum(
            peso for chave, peso, ativo in horizontes
            if ativo and not math.isnan(resultado[chave])
        )
        pontuacao_base = (
            sum(parcelas) / pesos_validos if pesos_validos else float("nan")
        )
        resultado["pontuacao_multitemporal_base"] = pontuacao_base
        resultado["pontuacao_multitemporal"] = pontuacao_base

    if meses >= 12:
        confianca = "ALTA"
    elif meses >= 6:
        confianca = "BOA"
    elif meses >= 3:
        confianca = "MODERADA"
    elif len(dias) >= 20:
        confianca = "MODERADA-BAIXA"
    elif len(dias) >= 2:
        confianca = "BAIXA"
    else:
        confianca = "MUITO BAIXA"
    return {
        "longo": usar_longo,
        "mensal": usar_mensal,
        "diario": usar_diario,
        "peso_disponivel": peso_disponivel,
        "confianca": confianca,
    }


def _moeda(valor: float) -> str:
    return f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")



def _imprimir_ranking_simples(titulo: str, resultados: list[dict], chave_nome: str = "estrategia") -> None:
    if not resultados:
        print(f"\n{titulo}: nenhuma combinação produziu resultado.")
        return
    colunas = (
        ("pos", "pos", False, 3),
        (chave_nome, "estratégia", False, 34),
        ("resultado", "resultado", True, 10),
        ("diaria", "diária", True, 7),
        ("dias_operados", "dias", False, 4),
        ("aproveitamento", "aproveit", True, 8),
        ("drawdown", "drawdown", True, 10),
        ("capital_minimo", "cap_min", True, 9),
        ("acumulado", "acumulado", True, 10),
    )
    largura_total = sum(c[3] for c in colunas) + len(colunas) - 1
    print(f"\n========== {titulo} ==========")
    print(" ".join(f"{rotulo:>{largura}}" for _, rotulo, _, largura in colunas))
    print("-" * largura_total)

    def formatar(linha, chave, moeda):
        valor = linha[chave]
        if moeda:
            return _moeda(valor)
        return str(valor)

    def ajustar(texto, largura):
        return texto if len(texto) <= largura else texto[:largura - 2] + ".."

    for indice, linha in enumerate(resultados):
        cor = (
            AZUL_CELESTE if linha.get("titular")
            else VERDE if indice == 0
            else (VERMELHO if indice == len(resultados) - 1 else "")
        )
        texto = " ".join(
            f"{ajustar(formatar(linha, chave, moeda), largura):>{largura}}"
            for chave, _, moeda, largura in colunas
        )
        print(f"{cor}{texto}{RESET}")
    print(
        f"{VERDE}Verde{RESET}=1º | {VERMELHO}Vermelho{RESET}=último | "
        f"{AZUL_CELESTE}Azul{RESET}=titular do slot"
    )
    print("Acumulado = resultado - abs(drawdown). Maior acumulado vence.")


def _imprimir_ranking_cruzado(resultados: list[dict], total_pregoes: int = 0) -> None:
    if not resultados:
        print("\n========== RANKING PRINCIPAL (ENTRADA × SAÍDA) ==========")
        print("Nenhuma combinação produziu resultado.")
        return
    colunas = (
        ("pos", "pos", False, 3),
        ("entrada", "entrada", False, 24),
        ("saida", "saída", False, 18),
        ("resultado", "resultado", True, 10),
        ("diaria", "diária", True, 7),
        ("dias_operados", "dias", False, 4),
        ("drawdown", "drawdown", True, 10),
        ("capital_minimo", "cap_min", True, 9),
        ("acumulado", "acumulado", True, 10),
    )
    largura_total = sum(c[3] for c in colunas) + len(colunas) - 1
    print(
        f"\n========== RANKING PRINCIPAL {VERSAO_CLASSIFICACAO} "
        f"(ENTRADA × SAÍDA | MOTOR {cfg.VERSAO}) =========="
    )
    print(f"Total de pregões avaliados: {total_pregoes}")
    print(" ".join(f"{rotulo:>{largura}}" for _, rotulo, _, largura in colunas))
    print("-" * largura_total)

    def formatar(linha, chave, moeda):
        valor = linha[chave]
        if moeda:
            return _moeda(valor)
        return str(valor)

    def ajustar(texto, largura):
        return texto if len(texto) <= largura else texto[:largura - 2] + ".."

    for indice, linha in enumerate(resultados):
        cor = VERDE if indice == 0 else (VERMELHO if indice == len(resultados) - 1 else "")
        texto = " ".join(
            f"{ajustar(formatar(linha, chave, moeda), largura):>{largura}}"
            for chave, _, moeda, largura in colunas
        )
        print(f"{cor}{texto}{RESET}")
    print("Acumulado = resultado - abs(drawdown). Maior acumulado vence.")


def _montar_resultado_par(
    nome_entrada: str,
    nome_saida: str,
    trades: list[dict],
    dias: list,
    entrada_titular: bool = False,
    saida_titular: bool = False,
) -> dict:
    resumo = _resumo(trades, len(dias))
    corte = dias[-min(cfg.LAB_DIAS_TESTE, max(1, len(dias) // 3))]
    longo = [t for t in trades if t["horario_rotulo"].date() < corte]
    curto = [t for t in trades if t["horario_rotulo"].date() >= corte]
    pf_longo = _profit_factor(longo)
    pf_curto = _profit_factor(curto)
    pf_minimo = (
        min(pf_longo, pf_curto)
        if not math.isnan(pf_longo) and not math.isnan(pf_curto)
        else float("nan")
    )
    score = _pontuacao(
        resumo["resultado"], resumo["capital_minimo"], pf_minimo, resumo["dias_operados"]
    )
    acumulado = resumo["resultado"] - abs(resumo["drawdown"])
    robustez = avaliar_robustez(trades, dias, cfg.LAB_JANELA_PREGOES, cfg.LAB_PASSO_JANELA)
    horizontes = avaliar_horizontes(trades, dias)
    return {
        "entrada": nome_entrada,
        "saida": nome_saida,
        "estrategia": nome_entrada,  # ranking só-entrada
        "estrategia_saida": nome_saida,  # ranking só-saída
        "titular": entrada_titular and saida_titular,
        "titular_entrada": entrada_titular,
        "titular_saida": saida_titular,
        **resumo,
        "pf_longo": pf_longo,
        "pf_curto": pf_curto,
        "acumulado": acumulado,
        "score_longo_raw": score,      # métrica multitemporal mantida para auditoria
        **horizontes,
        **robustez,
        "operacoes_n": len(trades),
    }


def _perguntar_modo_ranking() -> str:
    """
    Pergunta qual ranking rodar:
      E = só entradas (pareadas com saída titular)
      S = só saídas   (pareadas com entrada titular)
      C = cruzado     (todas as combinações — ranking principal)
      A = análise     (cruzado completo + planilha .xlsx: rankings e apuração
                       diária/mensal/anual por estratégia)
      D = diagnóstico (como cada entrada e cada saída se sai em cada cenário)
    """
    while True:
        resp = input(
            "Qual ranking deseja? Entradas (E), Saídas (S), Cruzada (C), Análise (A) "
            "ou Diagnóstico de cenários (D): "
        ).strip().upper()
        if resp in ("E", "S", "C", "A", "D"):
            return resp
        print(f"Resposta '{resp}' não reconhecida — digite exatamente E, S, C, A ou D.")


def _perguntar_cenario() -> str:
    """'' = todos os cenários; senão, o nome de um cenário de cenario.py."""
    while True:
        resp = input(
            "Cenário [Enter = todos | " + " | ".join(cen_mod.NOMES) + "]: "
        ).strip().lower()
        if resp in ("", "todos", "tudo"):
            return ""
        if resp in cen_mod.NOMES:
            return resp
        print(f"Cenário '{resp}' não existe. Opções: {', '.join(cen_mod.NOMES)}")


def _perguntar_simulacao() -> Optional[dict]:
    """None = usar o histórico real; dict = parâmetros do simulador."""
    resp = input("Fonte dos dados: Real (R) ou Simulada (S)? [Enter = R]: ").strip().upper()
    if resp not in ("S", "SIM", "SIMULADA"):
        return None
    modo = input("Simulação: Reamostragem de pregões reais (R) ou Regimes sintéticos (G)? [Enter = R]: "
                 ).strip().upper()
    n = input("Quantos pregões simular? [Enter = 120]: ").strip()
    sem = input("Semente (número) [Enter = aleatória]: ").strip()
    return {
        "modo": "regimes" if modo == "G" else "reamostragem",
        "dias": int(n) if n.isdigit() else 120,
        "semente": int(sem) if sem.isdigit() else None,
    }


# ---------------------------------------------------------------------------
# Modo A (Análise) — planilha .xlsx. Incorpora o antigo analise.py: mesma engine
# (executar_jogo), nenhuma lógica de backtest duplicada.
# ---------------------------------------------------------------------------
PASTA_ANALISES_PADRAO = r"D:\DAYTRADE\ANALISES"
_DIAS_SEMANA = ("segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo")
_FORMATO_MOEDA_XLSX = '"R$ "#,##0.00;[Red]"(R$ "#,##0.00\\);\\-'


def _resultado_por_dia(trades: list[dict]) -> dict:
    """{data: resultado líquido do dia} — a data é a do candle de sinal."""
    por_dia: dict = {}
    for t in trades:
        d = t["horario_rotulo"].date()
        por_dia[d] = por_dia.get(d, 0.0) + t["resultado_reais"]
    return por_dia


def _colunas_apuracao(rank_entrada: list[dict], rank_saida: list[dict],
                      rank_cruzado: list[dict]) -> list[tuple[str, str, dict]]:
    """Colunas das abas de apuração: (bloco, rótulo, {dia: resultado}).

    Entradas (pareadas com a saída titular) e Saídas (pareadas com a entrada
    titular) aparecem inteiras; o par titular×titular está nos dois blocos. O
    bloco Cruzadas traz só os pares sem titular, que não estão nos outros dois.
    Dentro de cada bloco, a ordem é a do ranking (maior acumulado primeiro)."""
    colunas = []
    for r in rank_entrada:
        colunas.append(("Entradas", r["entrada"] + (" (titular)" if r["titular"] else ""), r["_por_dia"]))
    for r in rank_saida:
        colunas.append(("Saídas", r["saida"] + (" (titular)" if r["titular"] else ""), r["_por_dia"]))
    for r in rank_cruzado:
        if r["titular_entrada"] or r["titular_saida"]:
            continue
        colunas.append(("Cruzadas", f"{r['entrada']} × {r['saida']}", r["_por_dia"]))
    return colunas


def _agrupar_periodo(por_dia: dict, chave) -> dict:
    agrupado: dict = {}
    for d, v in por_dia.items():
        k = chave(d)
        agrupado[k] = agrupado.get(k, 0.0) + v
    return agrupado


def _estilos_xlsx():
    from openpyxl.styles import Alignment, Font, PatternFill
    return {
        "Alignment": Alignment,
        "titulo": Font(name="Arial", bold=True, size=14),
        "cab_fonte": Font(name="Arial", bold=True, color="FFFFFFFF"),
        "cab_fill": PatternFill(start_color="FF17365D", end_color="FF17365D", fill_type="solid"),
        "bloco_fill": {
            "Entradas": PatternFill(start_color="FF2E75B6", end_color="FF2E75B6", fill_type="solid"),
            "Saídas": PatternFill(start_color="FF548235", end_color="FF548235", fill_type="solid"),
            "Cruzadas": PatternFill(start_color="FF7F6000", end_color="FF7F6000", fill_type="solid"),
        },
        "titular_fill": PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid"),
        "fonte": Font(name="Arial", size=10),
        "negrito": Font(name="Arial", size=10, bold=True),
    }


def _aba_ranking(wb, nome_aba: str, titulo: str, linhas: list[dict], par: bool) -> None:
    from openpyxl.utils import get_column_letter
    est = _estilos_xlsx()
    ws = wb.create_sheet(nome_aba)
    ws.cell(row=1, column=1, value=titulo).font = est["titulo"]
    ws.cell(row=2, column=1,
            value="Ordem: acumulado = resultado − |drawdown| (maior vence). Linha azul = titular.").font = est["fonte"]
    colunas = [("pos", "Pos", 6)]
    if par:
        colunas += [("entrada", "Entrada", 36), ("saida", "Saída", 30)]
    else:
        colunas += [("estrategia", "Estratégia", 38)]
    colunas += [
        ("resultado", "Resultado (R$)", 15), ("diaria", "Diária (R$)", 12),
        ("dias_operados", "Dias operados", 10), ("aproveitamento", "Aproveit. (R$/dia op.)", 14),
        ("drawdown", "Drawdown (R$)", 14), ("capital_minimo", "Capital mín. (R$)", 15),
        ("acumulado", "Acumulado (R$)", 15), ("operacoes", "Operações", 10), ("pf", "Profit factor", 10),
    ]
    moeda = {"resultado", "diaria", "aproveitamento", "drawdown", "capital_minimo", "acumulado"}
    for c, (_, rotulo, largura) in enumerate(colunas, 1):
        cel = ws.cell(row=4, column=c, value=rotulo)
        cel.font, cel.fill = est["cab_fonte"], est["cab_fill"]
        cel.alignment = est["Alignment"](horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = largura
    ws.row_dimensions[4].height = 32
    for i, linha in enumerate(linhas):
        for c, (chave, _, _) in enumerate(colunas, 1):
            valor = linha.get(chave)
            if isinstance(valor, float):
                valor = None if math.isnan(valor) or math.isinf(valor) else round(valor, 2)
            cel = ws.cell(row=5 + i, column=c, value=valor)
            cel.font = est["fonte"]
            if chave in moeda:
                cel.number_format = _FORMATO_MOEDA_XLSX
            if linha.get("titular"):
                cel.fill = est["titular_fill"]
    ws.freeze_panes = "A5"


def _aba_apuracao(wb, nome_aba: str, titulo: str, colunas: list, chaves: list,
                  rotulo_chave: str, chave_fn, rotulo_fn, formato_chave,
                  semana: bool = False) -> None:
    """Linhas = períodos (dia/mês/ano); colunas = estratégias. Última linha = total.

    chaves: períodos já ordenados; chave_fn(data) -> chave do período;
    rotulo_fn(chave) -> valor exibido na primeira coluna."""
    from openpyxl.utils import get_column_letter
    est = _estilos_xlsx()
    ws = wb.create_sheet(nome_aba)
    ws.cell(row=1, column=1, value=titulo).font = est["titulo"]
    ws.cell(row=2, column=1,
            value="Resultado líquido (R$, 1 contrato) por período; célula '-' = sem operação/zero.").font = est["fonte"]
    deslocamento = 2 if semana else 1
    cab = ws.cell(row=4, column=1, value=rotulo_chave)
    cab.font, cab.fill = est["cab_fonte"], est["cab_fill"]
    cab.alignment = est["Alignment"](horizontal="center", vertical="center")
    ws.column_dimensions["A"].width = 12
    if semana:
        c2 = ws.cell(row=4, column=2, value="Dia")
        c2.font, c2.fill = est["cab_fonte"], est["cab_fill"]
        c2.alignment = est["Alignment"](horizontal="center", vertical="center")
        ws.column_dimensions["B"].width = 9
    agrupados = []
    for c, (bloco, rotulo, por_dia) in enumerate(colunas, deslocamento + 1):
        bl = ws.cell(row=3, column=c, value=bloco)
        bl.font, bl.fill = est["cab_fonte"], est["bloco_fill"][bloco]
        bl.alignment = est["Alignment"](horizontal="center")
        cel = ws.cell(row=4, column=c, value=rotulo)
        cel.font, cel.fill = est["cab_fonte"], est["cab_fill"]
        cel.alignment = est["Alignment"](horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = 18
        agrupados.append(_agrupar_periodo(por_dia, chave_fn))
    ws.row_dimensions[4].height = 78
    totais = [0.0] * len(colunas)
    for i, chave in enumerate(chaves):
        linha = 5 + i
        k = ws.cell(row=linha, column=1, value=rotulo_fn(chave))
        k.font = est["fonte"]
        k.number_format = formato_chave
        if semana:
            ws.cell(row=linha, column=2, value=_DIAS_SEMANA[chave.weekday()]).font = est["fonte"]
        for j, agr in enumerate(agrupados):
            bruto = agr.get(chave, 0.0)
            totais[j] += bruto
            cel = ws.cell(row=linha, column=deslocamento + 1 + j, value=round(bruto, 2))
            cel.font = est["fonte"]
            cel.number_format = _FORMATO_MOEDA_XLSX
    linha_total = 5 + len(chaves)
    ws.cell(row=linha_total, column=1, value="Total").font = est["negrito"]
    for j, total in enumerate(totais):
        cel = ws.cell(row=linha_total, column=deslocamento + 1 + j, value=round(total, 2))
        cel.font = est["negrito"]
        cel.number_format = _FORMATO_MOEDA_XLSX
    ws.freeze_panes = ws.cell(row=5, column=deslocamento + 1)


def _aba_operacoes_titulares(wb, trades: list[dict], meta: dict) -> None:
    from openpyxl.utils import get_column_letter
    est = _estilos_xlsx()
    ws = wb.create_sheet("Operações titulares")
    ws.cell(row=1, column=1,
            value=f"Operações dos titulares ({meta['entrada_titular']} × {meta['saida_titular']})"
            ).font = est["titulo"]
    ws.cell(row=2, column=1, value=f"{len(trades)} operações | resultado líquido R$ "
            f"{sum(t['resultado_reais'] for t in trades):,.2f}").font = est["fonte"]
    cabecalho = ["Dia", "Data", "Hora de entrada", "Hora de saída", "Tipo da operação",
                 "Resultado (R$)", "Total acumulado (R$)", "Motivo"]
    for c, rotulo in enumerate(cabecalho, 1):
        cel = ws.cell(row=4, column=c, value=rotulo)
        cel.font, cel.fill = est["cab_fonte"], est["cab_fill"]
        cel.alignment = est["Alignment"](horizontal="center", vertical="center")
    acumulado = 0.0
    for i, t in enumerate(trades):
        linha = 5 + i
        data_op = t["horario_rotulo"].date()
        saida_fim = t["saida_dt"]
        saida_ini = saida_fim - timedelta(minutes=cfg.TIMEFRAME_MINUTOS)
        acumulado += t["resultado_reais"]
        motivo = t["motivo"] + (" (ambíguo)" if t.get("intrabar_ambiguo") else "")
        valores = [_DIAS_SEMANA[data_op.weekday()], data_op, t["horario_execucao"].strftime("%H:%M"),
                   f"{saida_ini:%H:%M}-{saida_fim:%H:%M}", t["lado"],
                   round(t["resultado_reais"], 2), round(acumulado, 2), motivo]
        for c, v in enumerate(valores, 1):
            cel = ws.cell(row=linha, column=c, value=v)
            cel.font = est["fonte"]
            if c == 2:
                cel.number_format = "dd/mm/yyyy"
            if c in (6, 7):
                cel.number_format = _FORMATO_MOEDA_XLSX
    for c, w in enumerate((10, 12, 15, 15, 17, 16, 20, 24), 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A5"


def _aba_metodologia(wb, meta: dict) -> None:
    est = _estilos_xlsx()
    ws = wb.create_sheet("Metodologia")
    ws.cell(row=1, column=1, value="Metodologia e critérios").font = est["titulo"]
    for c, rotulo in enumerate(("Item", "Critério aplicado"), 1):
        cel = ws.cell(row=3, column=c, value=rotulo)
        cel.font, cel.fill = est["cab_fonte"], est["cab_fill"]
    janelas = " e ".join(f"{ini}–{fim}" for ini, fim in cfg.JANELAS_BLOQUEADAS)
    itens = [
        ("Ordenação dos rankings", "Acumulado = resultado − |drawdown|. Sem filtro de elegibilidade, sem pontos por dia e sem portão estatístico."),
        ("Entradas", f"Cada entrada pareada com a saída titular ({meta['saida_titular']}.py)."),
        ("Saídas", f"Cada saída pareada com a entrada titular ({meta['entrada_titular']}.py)."),
        ("Cruzadas", "Todas as combinações entrada × saída. Nas abas de apuração, o bloco Cruzadas omite os pares que já aparecem nos blocos Entradas/Saídas (os que têm titular)."),
        ("Apuração diária / mensal / anual", "Soma do resultado líquido por dia do candle de sinal, por mês e por ano. Dias sem operação valem zero."),
        ("Fonte e período", meta["fonte"]),
        ("Histórico utilizado",
         f"{meta['arquivo']}, de {meta['historico_inicio']:%d/%m/%Y} a {meta['historico_fim']:%d/%m/%Y} "
         f"({meta['n_candles']} candles, timeframe {cfg.TIMEFRAME_MINUTOS} min). Candles fora do período "
         "só aquecem os indicadores."),
        ("Motor", f"Motor oficial do laboratório (executar_jogo): stop e alvo vêm da saída desde a abertura; custo de R$ "
         f"{cfg.CUSTO_TOTAL_ESTIMADO_POR_OPERACAO_REAIS:.2f} por operação; R$ {cfg.VALOR_PONTO_REAIS:.2f} por ponto de WIN."),
        ("Limites", f"Máximo de {cfg.MAX_PERDAS_DIA} perdas por dia; janela(s) {janelas} bloqueada(s)."),
        ("Hora de entrada", f"Fechamento do candle de sinal (rótulo + {cfg.TIMEFRAME_MINUTOS} minutos)."),
        ("Hora de saída", "Intervalo do candle em que stop, alvo, saída do cartucho ou encerramento ocorreu; o OHLC não informa o segundo exato."),
        ("Limitação", "Backtest histórico não garante resultado futuro."),
    ]
    for i, (campo, valor) in enumerate(itens, start=4):
        ws.cell(row=i, column=1, value=campo).font = est["negrito"]
        cel = ws.cell(row=i, column=2, value=valor)
        cel.font = est["fonte"]
        cel.alignment = est["Alignment"](wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110


def gravar_planilha_analise(
    caminho: Path, *, rank_entrada: list[dict], rank_saida: list[dict],
    rank_cruzado: list[dict], dias: list, trades_titular: list[dict], meta: dict,
) -> Path:
    """Grava o .xlsx do modo A. Abas: Entradas, Saídas, Cruzadas (rankings),
    Apuração diária, mensal e anual (período nas linhas, estratégia nas colunas),
    Operações titulares e Metodologia."""
    try:
        from openpyxl import Workbook
    except ImportError:
        raise SystemExit("Pacote 'openpyxl' não encontrado. Instale com:\n    pip install openpyxl")
    wb = Workbook()
    wb.remove(wb.active)
    sufixo = f" — {meta['fonte']}" if meta.get("fonte") else ""
    _aba_ranking(wb, "Entradas", f"Ranking de entradas (saída titular: {meta['saida_titular']}){sufixo}",
                 rank_entrada, par=False)
    _aba_ranking(wb, "Saídas", f"Ranking de saídas (entrada titular: {meta['entrada_titular']}){sufixo}",
                 rank_saida, par=False)
    _aba_ranking(wb, "Cruzadas", f"Ranking cruzado entrada × saída{sufixo}", rank_cruzado, par=True)

    colunas = _colunas_apuracao(rank_entrada, rank_saida, rank_cruzado)
    dias = sorted(dias)
    meses = sorted({(d.year, d.month) for d in dias})
    anos = sorted({d.year for d in dias})
    _aba_apuracao(wb, "Apuração diária", f"Apuração diária{sufixo}", colunas, dias,
                  "Data", lambda d: d, lambda k: k, "dd/mm/yyyy", semana=True)
    _aba_apuracao(wb, "Apuração mensal", f"Apuração mensal{sufixo}", colunas, meses,
                  "Mês", lambda d: (d.year, d.month), lambda k: f"{k[1]:02d}/{k[0]}", "@")
    _aba_apuracao(wb, "Apuração anual", f"Apuração anual{sufixo}", colunas, anos,
                  "Ano", lambda d: d.year, lambda k: k, "0")
    _aba_operacoes_titulares(wb, trades_titular, meta)
    _aba_metodologia(wb, meta)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return caminho


# ---------------------------------------------------------------------------
# Modo D (Diagnóstico de cenários) — incorpora o antigo diagnostico_cenarios.py.
# Só leitura: não altera motor, cartuchos nem ranking. Cada operação é marcada
# com o cenário (cenario.py) do candle que gerou o sinal; depois agrega por
# estratégia x cenário. Com dados simulados em 'regimes', mede também o quanto
# o reconhecedor de cenário acerta contra o gabarito do simulador.
DIAG_MIN_OPS_CONFIAVEL = 20     # abaixo disso a célula é só indício


def diagnosticar_cenarios(candles, rows, dias, entradas, saidas, entrada_titular, saida_titular) -> list[dict]:
    """Cada entrada x saida titular e cada saida x entrada titular, agregadas por cenario.
    Linhas: tipo ('entrada'|'saida'), estrategia, cenario, ops, resultado, vitorias."""
    row_por_hora = {c.horario: r for c, r in zip(candles, rows)}
    saidas_ok = list(saidas)       # ja validadas por executar() (contrato de saida)
    pares = [("entrada", e.nome, e, saida_titular) for e in entradas]
    pares += [("saida", s.nome, entrada_titular, s) for s in saidas_ok]

    celulas = defaultdict(lambda: [0, 0.0, 0])
    for tipo, nome, ent, sai in pares:
        try:
            trades = executar_jogo(candles, rows, ent, dias, avaliar_saida=sai.avaliar_saida)
        except (KeyError, TypeError, ValueError, AttributeError):
            continue
        for t in trades:
            cen = cen_mod.classificar(row_por_hora.get(t["horario_rotulo"]))
            c = celulas[(tipo, nome, cen)]
            c[0] += 1
            c[1] += t["resultado_reais"]
            c[2] += 1 if t["resultado_reais"] > 0 else 0
    return [
        {"tipo": k[0], "estrategia": k[1], "cenario": k[2], "ops": v[0],
         "resultado": round(v[1], 2), "vitorias": v[2]}
        for k, v in sorted(celulas.items())
    ]


def distribuicao_cenarios(rows) -> dict:
    cont = defaultdict(int)
    for r in rows:
        if r is not None:
            cont[cen_mod.classificar(r)] += 1
    return dict(cont)


# Regime verdadeiro do simulador -> cenários aceitos como acerto.
DIAG_ACERTOS = {
    "tendencia_alta": {"tendencia", "esticado"},
    "tendencia_baixa": {"tendencia", "esticado"},
    "lateral": {"lateral"},
    "volatil": {"volatil", "esticado"},
}


def matriz_confusao_cenarios(candles, rows, rotulos: dict) -> dict:
    """Gabarito (regime do simulador) x cenário reconhecido. Candles de abertura
    e fim de tarde são separados do resto, pois dependem do relógio, não do regime."""
    m = defaultdict(lambda: defaultdict(int))
    for c, r in zip(candles, rows):
        if r is None or c.horario not in rotulos:
            continue
        m[rotulos[c.horario]][cen_mod.classificar(r)] += 1
    return {k: dict(v) for k, v in m.items()}


def acerto_geral_cenarios(matriz: dict) -> float:
    ok = tot = 0
    for regime, cols in matriz.items():
        for cen, n in cols.items():
            if cen in ("abertura", "fim_de_tarde", cen_mod.INDEFINIDO):
                continue
            tot += n
            if cen in DIAG_ACERTOS.get(regime, set()):
                ok += n
    return ok / tot if tot else float("nan")


_DIAG_LARGURA_NOME = 24
_DIAG_LARGURA_CEL = 11           # 24 + 7*11 = 101 colunas: cabe ate num console estreito
_DIAG_IAS = ("chatgpt", "claude", "deepseek", "gemini", "grok", "manus")


def nomes_curtos(nomes) -> dict:
    """Nome para exibir: sem o prefixo entrada_/saida_ (o titulo da tabela ja diz)
    e sem o nome da IA quando sobra uma descricao (macd_estocastico_claude_v1 ->
    macd_estocastico_v1). Cartuchos antigos, cujo nome e so a IA, ficam como estao.
    Se dois nomes curtos colidirem, os dois voltam ao nome sem prefixo."""
    base = {n: re.sub(r"^(entrada|saida)_", "", n) for n in nomes}
    curto = {}
    for n, b in base.items():
        partes = b.split("_")
        sem_ia = [x for x in partes if x.lower() not in _DIAG_IAS]
        descricao = [x for x in sem_ia if not re.fullmatch(r"v?\d+", x, re.I)]
        curto[n] = "_".join(sem_ia) if descricao else b     # sem descricao, a IA e o proprio nome
    repetidos = {c for c in curto.values() if list(curto.values()).count(c) > 1}
    return {n: (base[n] if curto[n] in repetidos else curto[n]) for n in nomes}


def imprimir_distribuicao_cenarios(dist: dict) -> None:
    total = sum(dist.values()) or 1
    print("\nCandles por cenario:")
    for nome in cen_mod.NOMES:
        n = dist.get(nome, 0)
        print(f"  {nome:14s}{n:>7d}  {100 * n / total:>4.0f}%")


def imprimir_confusao_cenarios(matriz: dict) -> None:
    cols = list(cen_mod.NOMES)
    print("\nGabarito do simulador (linhas) x cenario reconhecido (colunas), % da linha:")
    print(f"  {'gabarito':17s}" + "".join(f"{c[:9]:>10s}" for c in cols))
    for regime, v in matriz.items():
        tot = sum(v.values()) or 1
        print(f"  {regime:17s}" + "".join(f"{100 * v.get(c, 0) / tot:>9.0f}%" for c in cols))


def imprimir_diagnostico(linhas: list[dict], titulo: str = "") -> None:
    if titulo:
        print(f"\n===== {titulo} =====")
    rotulos = {"abertura": "abertura", "fim_de_tarde": "fim_tarde", "volatil": "volatil",
               "esticado": "esticado", "tendencia": "tendencia", "lateral": "lateral",
               cen_mod.INDEFINIDO: "indef."}
    for tipo in ("entrada", "saida"):
        sub = [l for l in linhas if l["tipo"] == tipo]
        if not sub:
            continue
        nomes = sorted({l["estrategia"] for l in sub})
        curtos = nomes_curtos(nomes)
        cols = list(cen_mod.NOMES)
        print(f"\n-- {tipo.upper()}: R$/ops por cenario; * = menos de {DIAG_MIN_OPS_CONFIAVEL} ops --")
        print(f"{'':{_DIAG_LARGURA_NOME}s}" + "".join(f"{rotulos[c]:>{_DIAG_LARGURA_CEL}s}" for c in cols))
        for n in nomes:
            partes = []
            for c in cols:
                cel = next((l for l in sub if l["estrategia"] == n and l["cenario"] == c), None)
                if cel is None:
                    partes.append(f"{'-':>{_DIAG_LARGURA_CEL}s}")
                else:
                    marca = "*" if cel["ops"] < DIAG_MIN_OPS_CONFIAVEL else " "
                    partes.append(f"{cel['resultado']:.0f}/{cel['ops']}{marca}".rjust(_DIAG_LARGURA_CEL))
            print(f"{curtos[n][:_DIAG_LARGURA_NOME]:{_DIAG_LARGURA_NOME}s}" + "".join(partes))


DIAG_MIN_OPS_METADE = 10        # operações mínimas por metade para a célula contar


def consistencia_metades(linhas_a: list[dict], linhas_b: list[dict], nome_titular_entrada: str,
                 nome_titular_saida: str) -> list[dict]:
    """Compara duas metades (do período que o usuário escolheu) e devolve as células
    em que a estratégia fica melhor (ou pior) que a titular NO MESMO SENTIDO nas
    duas metades. Só conta célula com >= DIAG_MIN_OPS_METADE operações em ambas as
    metades, na estratégia e na titular. Ordena pelo menor efeito entre as metades.
    Não define período nenhum: as metades vêm dos dias que foram apurados."""
    def indexa(linhas):
        return {(l["tipo"], l["estrategia"], l["cenario"]): l for l in linhas}

    a, b = indexa(linhas_a), indexa(linhas_b)
    saida = []
    for (tipo, nome, cen), la in a.items():
        titular = nome_titular_entrada if tipo == "entrada" else nome_titular_saida
        if nome == titular or (tipo, nome, cen) not in b:
            continue
        ta, tb = a.get((tipo, titular, cen)), b.get((tipo, titular, cen))
        lb = b[(tipo, nome, cen)]
        if None in (ta, tb):
            continue
        if min(la["ops"], lb["ops"], ta["ops"], tb["ops"]) < DIAG_MIN_OPS_METADE:
            continue
        da, db = la["resultado"] - ta["resultado"], lb["resultado"] - tb["resultado"]
        if da * db > 0:
            saida.append({"tipo": tipo, "estrategia": nome, "cenario": cen,
                          "delta_metade_1": round(da, 2), "delta_metade_2": round(db, 2),
                          "efeito_minimo": round(min(abs(da), abs(db)) * (1 if da > 0 else -1), 2)})
    return sorted(saida, key=lambda x: -abs(x["efeito_minimo"]))


def gravar_csv_diagnostico(linhas: list[dict], caminho: Path) -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["tipo", "estrategia", "cenario", "ops", "resultado", "vitorias"],
                           delimiter=";")
        w.writeheader()
        w.writerows(linhas)
    return caminho


def _executar_diagnostico(candles, rows, dias, dias_avaliacao, entradas, saidas,
                          entrada_titular, saida_titular, incompativeis, fonte_arquivo,
                          rotulos_sim, metades) -> dict:
    """Modo D: imprime distribuição de cenários, (se simulado em regimes) a matriz
    de confusão do reconhecedor, a tabela estratégia × cenário e grava o CSV."""
    print("[ETAPA 2/2] Diagnóstico por cenário (cada entrada × saída titular e cada saída × entrada titular)...")
    imprimir_distribuicao_cenarios(distribuicao_cenarios(rows))
    if rotulos_sim:
        matriz = matriz_confusao_cenarios(candles, rows, rotulos_sim)
        imprimir_confusao_cenarios(matriz)
        print(f"Acerto (fora abertura/fim de tarde): {acerto_geral_cenarios(matriz):.1%}")
    linhas = diagnosticar_cenarios(
        candles, rows, dias_avaliacao, entradas, saidas, entrada_titular, saida_titular)
    imprimir_diagnostico(linhas, f"DIAGNÓSTICO DE CENÁRIOS {VERSAO_CLASSIFICACAO} (motor {cfg.VERSAO})")

    if metades is None:
        metades = input("Mostrar o que se repete nas duas metades do período? (S/N) [Enter = N]: "
                        ).strip().upper() in ("S", "SIM")
    repetem = []
    if metades and len(dias) >= 2:
        meio = len(dias) // 2
        d1, d2 = set(dias[:meio]), set(dias[meio:])
        l1 = diagnosticar_cenarios(candles, rows, d1, entradas, saidas, entrada_titular, saida_titular)
        l2 = diagnosticar_cenarios(candles, rows, d2, entradas, saidas, entrada_titular, saida_titular)
        repetem = consistencia_metades(l1, l2, entrada_titular.nome, saida_titular.nome)
        print("\n-- REPETE NAS DUAS METADES, em R$ vs a titular (min "
              f"{DIAG_MIN_OPS_METADE} ops por metade) --")
        print(f"metade 1: {dias[0]} a {dias[meio - 1]}")
        print(f"metade 2: {dias[meio]} a {dias[-1]}")
        curtos = nomes_curtos({r["estrategia"] for r in repetem})
        print(f"{'':8s}{'':26s}{'cenario':13s}{'metade 1':>10s}{'metade 2':>10s}")
        for r in repetem[:25]:
            print(f"{r['tipo']:8s}{curtos[r['estrategia']][:24]:26s}{r['cenario']:13s}"
                  f"{r['delta_metade_1']:>10.0f}{r['delta_metade_2']:>10.0f}")
        if not repetem:
            print("nenhuma célula se repete nas duas metades")
        print("Muitas células são testadas: um padrão que repete ainda pode ser acaso; trate como hipótese.")

    csv_diag = gravar_csv_diagnostico(
        linhas,
        Path(cfg.PASTA_LOGS_AUDITORIA) / f"diagnostico_cenarios_{VERSAO_CLASSIFICACAO.lower()}"
                                         f"_motor_{cfg.VERSAO.lower()}{fonte_arquivo}.csv",
    )
    print(f"\nRelatórios:\n  {csv_diag.resolve()}")
    print("Limitação: OHLC 15min não revela ordem intrabar nem preço exato das 18:20.")
    if incompativeis:
        print("\nCartuchos/combinações INCOMPATIVEIS:")
        for nome, erro in incompativeis:
            print(f"- {nome} (INCOMPATIVEL): {erro}")
    return {"modo": "D", "diagnostico": linhas, "repetem_nas_metades": repetem, "csv": csv_diag}


def executar(
    caminho_csv: Optional[str] = None,
    *,
    modo: Optional[str] = None,
    periodo: Optional[str] = None,
    simulacao: Optional[dict] = None,
    cenario: Optional[str] = None,
    saida_dir: Optional[str] = None,
    metades: Optional[bool] = None,
):
    """modo: 'E'/'S'/'C'/'A'/'D' (None = pergunta; A = análise: cruzado completo +
    planilha .xlsx em saida_dir, padrão D:\\DAYTRADE\\ANALISES; D = diagnóstico por
    cenário, só leitura, com todos os cenários; metades=True mostra o que se repete
    nas duas metades do período, None = pergunta). periodo: texto aceito por
    interpretar_periodo (None ou '' = tudo, o conteúdo do arquivo). simulacao: dict do
    simulador (modo/dias/semente/escala_vol/espelhar); None = pergunta a fonte;
    False = força o histórico real sem perguntar. cenario: nome de um cenário
    (cenario.py) para apurar o campeonato SÓ naquela situação; '' = todos;
    None = pergunta."""
    caminho = Path(caminho_csv or cfg.CAMINHO_HISTORICO_INICIAL)
    if not caminho.exists():
        # V462: o Profit passou a gravar o export com datas no nome
        # (WINFUT_F_0_15min_01-01-2026_02-10-2026.csv). Tenta o CSV mais recente
        # da mesma pasta (mesmo ativo/timeframe) antes de perguntar.
        from historico_csv import resolver_csv_historico
        achado = resolver_csv_historico(str(caminho), cfg.TIMEFRAME_MINUTOS)
        if Path(achado).exists():
            print(f"[HISTÓRICO] '{caminho.name}' não existe; usando '{Path(achado).name}'.")
            caminho = Path(achado)
    if not caminho.exists():
        informado = input(
            f"CSV não encontrado em '{caminho}'. Informe o caminho completo: "
        ).strip().strip('"')
        caminho = Path(informado)

    if modo is None:
        modo = _perguntar_modo_ranking()
    modo = modo.upper()
    if modo not in ("E", "S", "C", "A", "D"):
        raise ValueError(f"modo '{modo}' inválido; use E, S, C, A ou D")
    rotulos_modo = {
        "E": "ENTRADA (pareado com saída titular)",
        "S": "SAÍDA (pareado com entrada titular)",
        "C": "CRUZADO (todas as combinações entrada × saída)",
        "A": "ANÁLISE (cruzado completo + planilha com rankings e apuração diária/mensal/anual)",
        "D": "DIAGNÓSTICO DE CENÁRIOS (R$/operações por estratégia × cenário; só leitura)",
    }
    print(f"[MODO] Ranking escolhido: {modo} — {rotulos_modo[modo]}")

    candles_reais = carregar_csv(caminho)
    if simulacao is None:
        simulacao = _perguntar_simulacao()
    fonte_rotulo = ""      # aparece no título do ranking e no nome do relatório
    fonte_arquivo = ""
    if simulacao:
        import simulador_mercado as sim
        candles_avaliacao, rotulos_sim, semente_usada = sim.gerar_candles(
            candles_reais,
            n_dias=int(simulacao.get("dias", 120)),
            semente=simulacao.get("semente"),
            modo=simulacao.get("modo", "reamostragem"),
            escala_vol=float(simulacao.get("escala_vol", 1.0)),
            espelhar=bool(simulacao.get("espelhar", True)),
        )
        tag = f"{simulacao.get('modo', 'reamostragem')}_s{semente_usada}_{simulacao.get('dias', 120)}d"
        pasta_sim = Path(cfg.PASTA_LOGS_AUDITORIA)
        csv_sim = sim.salvar_csv(candles_avaliacao, pasta_sim / f"simulacao_{tag}.csv")
        sim.salvar_rotulos(rotulos_sim, pasta_sim / f"simulacao_{tag}_rotulos.csv")
        print(f"[SIMULADOR] {simulacao.get('modo', 'reamostragem')} | semente {semente_usada} | "
              f"{simulacao.get('dias', 120)} pregões FICTÍCIOS | gravado em {csv_sim.resolve()}")
        print("[SIMULADOR] Dados fictícios: servem para estressar e comparar, não para estimar lucro.")
        fonte_rotulo = f" | SIMULADO {tag}"
        fonte_arquivo = f"_sim_{tag}"
        dias = sorted({c.horario.date() for c in candles_avaliacao})
        dias_avaliacao = set(dias)
        candles, quantidade_aquecimento = candles_avaliacao, 0
    else:
        todos_dias = sorted({c.horario.date() for c in candles_reais})
        # O período apurado é o do conteúdo do arquivo. Só um --periodo explícito
        # na linha de comando recorta; nunca há pergunta.
        dias, rotulo_periodo = interpretar_periodo(periodo or "", todos_dias)
        dias_avaliacao = set(dias)
        if rotulo_periodo != "tudo":
            fonte_rotulo = f" | período {rotulo_periodo}"
            fonte_arquivo = f"_{rotulo_periodo}"
        candles_avaliacao = [c for c in candles_reais if c.horario.date() in dias_avaliacao]
        # Os candles fora do período continuam disponíveis só como aquecimento.
        candles, quantidade_aquecimento = preparar_aquecimento(candles_reais)

    if modo == "D":
        cenario = ""        # o diagnóstico já separa por cenário: nunca recorta nem pergunta
    if cenario is None:
        cenario = _perguntar_cenario()
    cenario = (cenario or "").strip().lower()
    if cenario in ("todos", "tudo"):
        cenario = ""
    if cenario and cenario not in cen_mod.NOMES:
        raise ValueError(f"cenário '{cenario}' não existe; opções: {', '.join(cen_mod.NOMES)}")
    if cenario:
        fonte_rotulo += f" | cenário {cenario}"
        fonte_arquivo += f"_cen_{cenario}"
        print(f"[CENÁRIO] Campeonato só no cenário '{cenario}': cada entrada só pode disparar "
              "em candles desse cenário.")

    entradas, falhas_e = descobrir_entradas()
    saidas, falhas_s = descobrir_saidas()
    if not entradas:
        raise RuntimeError("Nenhuma entrada compatível foi encontrada.")
    if not saidas:
        raise RuntimeError("Nenhuma saída compatível foi encontrada.")

    entrada_titular = next(e for e in entradas if e.titular)
    saida_titular_carregada = next(s for s in saidas if s.titular)

    print("=" * 100)
    print(
        f"CLASSIFICAÇÃO [{VERSAO_CLASSIFICACAO}] — "
        f"ROBONILDO [{cfg.VERSAO}] ({rotulos_modo[modo]})"
    )
    print(
        f"Histórico: {candles_avaliacao[0].horario} até {candles_avaliacao[-1].horario} | "
        f"{len(candles_avaliacao)} candles | {len(dias)} pregões"
    )
    if quantidade_aquecimento:
        print(f"Aquecimento externo: {quantidade_aquecimento} candles (fora da pontuação)")
    print(f"Entrada titular: {entrada_titular.nome}.py")
    print(f"Saída titular  : {saida_titular_carregada.nome}.py")
    print(f"Entradas no ranking: {len(entradas)} | Saídas no ranking: {len(saidas)}")
    print("=" * 100)

    rows = preparar_rows(candles, dias_avaliacao)

    incompatíveis = list(falhas_e) + list(falhas_s)
    saidas_compativeis = []
    for sai in saidas:
        compativel, motivo = validar_contrato_saida(sai, rows)
        if compativel:
            saidas_compativeis.append(sai)
        else:
            incompatíveis.append((sai.nome, motivo))

    if saida_titular_carregada not in saidas_compativeis:
        motivo = next(
            (erro for nome, erro in incompatíveis if nome == saida_titular_carregada.nome),
            "contrato V445 inválido",
        )
        raise RuntimeError(
            f"A saída titular {saida_titular_carregada.nome}.py está INCOMPATIVEL: {motivo}"
        )

    saidas = saidas_compativeis
    saida_titular = saida_titular_carregada

    if modo == "D":
        return _executar_diagnostico(
            candles, rows, dias, dias_avaliacao, entradas, saidas, entrada_titular,
            saida_titular, incompatíveis, fonte_arquivo,
            rotulos_sim if simulacao else None, metades,
        )

    # Só calcula o que o modo pede — evita produto cartesiano desnecessário
    # e impede que uma saída sem proteção inicial produza números enganosos.
    if modo == "E":
        pares_a_rodar = [(ent, saida_titular) for ent in entradas]
    elif modo == "S":
        pares_a_rodar = [(entrada_titular, sai) for sai in saidas]
    else:  # C e A: todas as combinações (E e S do modo A são fatias do cruzado)
        pares_a_rodar = [(ent, sai) for ent in entradas for sai in saidas]

    pares_ok = []  # list of dict resultados
    trades_titular: list[dict] = []

    total = len(pares_a_rodar)
    n = 0
    for ent, sai in pares_a_rodar:
        n += 1
        rotulo = f"{ent.nome} × {sai.nome}"
        _progresso("COMBINAÇÕES", n - 1, total, rotulo)
        try:
            trades = executar_jogo(
                candles, rows, ent, dias_avaliacao, avaliar_saida=sai.avaliar_saida,
                cenario=cenario or None,
            )
        except (KeyError, TypeError, ValueError, AttributeError) as erro:
            incompatíveis.append((rotulo, f"{type(erro).__name__}: {erro}"))
            continue
        res = _montar_resultado_par(
            ent.nome, sai.nome, trades, dias,
            entrada_titular=ent.titular, saida_titular=sai.titular,
        )
        res["_por_dia"] = _resultado_por_dia(trades)    # modo A; não vai para o CSV
        if ent.titular and sai.titular:
            trades_titular = sorted(trades, key=lambda t: t["horario_execucao"])
        pares_ok.append(res)
    _progresso("COMBINAÇÕES", total, total)

    rank_entrada: list[dict] = []
    rank_saida: list[dict] = []
    rank_cruzado: list[dict] = []
    disponibilidade = {"longo": False, "mensal": False, "diario": False, "confianca": "—"}

    # ----- Ranking ENTRADA (modo E) -----
    if modo in ("E", "A"):
        for r in pares_ok:
            if not r["titular_saida"]:
                continue        # só entradas pareadas com a saída titular
            item = dict(r)
            item["estrategia"] = r["entrada"]
            item["titular"] = r["titular_entrada"]
            rank_entrada.append(item)
        disponibilidade = aplicar_pontuacao_multitemporal(rank_entrada, dias)
        for r in rank_entrada:
            if not disponibilidade["longo"]:
                r["nota_longo"] = float("nan")
            if not disponibilidade["mensal"]:
                r["nota_mensal"] = float("nan")
        rank_entrada.sort(
            key=lambda r: (
                r["acumulado"],
                r["resultado"],
                r["drawdown"],
                r["dias_operados"],
                -math.inf if math.isnan(r.get("pontuacao_robusta", float("nan"))) else r["pontuacao_robusta"],
            ),
            reverse=True,
        )
        for i, r in enumerate(rank_entrada, 1):
            r["pos"] = i
        titulo_entrada = (
            f"RANKING ENTRADA {VERSAO_CLASSIFICACAO} (motor {cfg.VERSAO}; "
            f"saída titular: {saida_titular.nome}){fonte_rotulo}"
        )
        with _capturar_e_imprimir() as _buf_entrada:
            _imprimir_ranking_simples(
                titulo_entrada,
                rank_entrada,
                chave_nome="estrategia",
            )
        _registrar_historico(titulo_entrada, _buf_entrada.getvalue())

    # ----- Ranking SAÍDA (modo S) -----
    if modo in ("S", "A"):
        for r in pares_ok:
            if not r["titular_entrada"]:
                continue        # só saídas pareadas com a entrada titular
            item = dict(r)
            item["estrategia"] = r["saida"]
            item["titular"] = r["titular_saida"]
            rank_saida.append(item)
        disponibilidade = aplicar_pontuacao_multitemporal(rank_saida, dias)
        for r in rank_saida:
            if not disponibilidade["longo"]:
                r["nota_longo"] = float("nan")
            if not disponibilidade["mensal"]:
                r["nota_mensal"] = float("nan")
        rank_saida.sort(
            key=lambda r: (
                r["acumulado"],
                r["resultado"],
                r["drawdown"],
                r["dias_operados"],
                -math.inf if math.isnan(r.get("pontuacao_robusta", float("nan"))) else r["pontuacao_robusta"],
            ),
            reverse=True,
        )
        for i, r in enumerate(rank_saida, 1):
            r["pos"] = i
        titulo_saida = (
            f"RANKING SAÍDA {VERSAO_CLASSIFICACAO} (motor {cfg.VERSAO}; "
            f"entrada titular: {entrada_titular.nome}){fonte_rotulo}"
        )
        with _capturar_e_imprimir() as _buf_saida:
            _imprimir_ranking_simples(
                titulo_saida,
                rank_saida,
                chave_nome="estrategia",
            )
        _registrar_historico(titulo_saida, _buf_saida.getvalue())

    # ----- Ranking CRUZADO (modo C) -----
    if modo in ("C", "A"):
        rank_cruzado = [dict(r) for r in pares_ok]
        rank_cruzado.sort(
            key=lambda r: (
                r["acumulado"],
                r["resultado"],
                r["drawdown"],
                r["dias_operados"],
            ),
            reverse=True,
        )
        for i, r in enumerate(rank_cruzado, 1):
            r["pos"] = i
        titulo_cruzado = f"RANKING CRUZADO {VERSAO_CLASSIFICACAO} (motor {cfg.VERSAO}){fonte_rotulo}"
        with _capturar_e_imprimir() as _buf_cruzado:
            _imprimir_ranking_cruzado(rank_cruzado, total_pregoes=len(dias))
        _registrar_historico(titulo_cruzado, _buf_cruzado.getvalue())

    if modo in ("E", "S", "A"):
        ativos = [n for n in ("longo", "mensal", "diario") if disponibilidade.get(n)]
        rotulo_h = {"E": "entrada", "S": "saída", "A": "análise"}[modo]
        print(f"\nHorizontes (ranking {rotulo_h}): {', '.join(ativos) or '—'} | "
              f"confiança: {disponibilidade['confianca']}")
        if disponibilidade["confianca"] in {"MUITO BAIXA", "BAIXA"}:
            print("Classificação provisória: histórico insuficiente para promover titular.")

    pasta_logs = Path(cfg.PASTA_LOGS_AUDITORIA)
    pasta_logs.mkdir(parents=True, exist_ok=True)

    def gravar(nome, linhas):
        if not linhas:
            return None
        saida = pasta_logs / (
            f"{nome}_{VERSAO_CLASSIFICACAO.lower()}_motor_{cfg.VERSAO.lower()}{fonte_arquivo}.csv"
        )
        campos = [k for k in linhas[0].keys() if not k.startswith("_")]
        with saida.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=campos, delimiter=";", extrasaction="ignore")
            w.writeheader()
            w.writerows(linhas)
        return saida

    relatorios = []
    if modo in ("E", "A"):
        relatorios.append(gravar("classificacao_entrada", rank_entrada))
    if modo in ("S", "A"):
        relatorios.append(gravar("classificacao_saida", rank_saida))
    if modo in ("C", "A"):
        relatorios.append(gravar("classificacao_cruzada", rank_cruzado))
    planilha = None
    if modo == "A":
        pasta_xlsx = Path(saida_dir or PASTA_ANALISES_PADRAO)
        nome_xlsx = re.sub(r"\s+", "_", f"analise_{dias[0]:%Y%m%d}_{dias[-1]:%Y%m%d}{fonte_arquivo}.xlsx")
        meta = {
            "entrada_titular": entrada_titular.nome,
            "saida_titular": saida_titular.nome,
            "fonte": (fonte_rotulo.strip(" |") or "histórico completo") ,
            "arquivo": caminho.name,
            "historico_inicio": candles_avaliacao[0].horario,
            "historico_fim": candles_avaliacao[-1].horario,
            "n_candles": len(candles_avaliacao),
        }
        try:
            planilha = gravar_planilha_analise(
                pasta_xlsx / nome_xlsx, rank_entrada=rank_entrada, rank_saida=rank_saida,
                rank_cruzado=rank_cruzado, dias=dias, trades_titular=trades_titular, meta=meta,
            )
        except OSError as erro:
            alternativa = Path(cfg.PASTA_LOGS_AUDITORIA) / nome_xlsx
            print(f"[AVISO] Não consegui gravar em '{pasta_xlsx}' ({erro}); gravando em '{alternativa.parent}'.")
            planilha = gravar_planilha_analise(
                alternativa, rank_entrada=rank_entrada, rank_saida=rank_saida,
                rank_cruzado=rank_cruzado, dias=dias, trades_titular=trades_titular, meta=meta,
            )
        relatorios.append(planilha)
    print("\nRelatórios:")
    for p in relatorios:
        if p:
            print(f"  {p.resolve()}")
    print("Limitação: OHLC 15min não revela ordem intrabar nem preço exato das 18:20.")
    if incompatíveis:
        print("\nCartuchos/combinações INCOMPATIVEIS:")
        for nome, erro in incompatíveis:
            print(f"- {nome} (INCOMPATIVEL): {erro}")
    return {
        "modo": modo,
        "entrada": rank_entrada,
        "saida": rank_saida,
        "cruzado": rank_cruzado,
        "planilha": planilha,
    }


def _ler_argumentos(argv=None):
    import argparse
    ap = argparse.ArgumentParser(
        description="Classificação de estratégias. Sem argumentos, pergunta tudo (comportamento original).")
    ap.add_argument("csv", nargs="?", help="histórico CSV (padrão: o configurado/mais recente)")
    ap.add_argument("--modo", choices=["E", "S", "C", "A", "D", "e", "s", "c", "a", "d"],
                    help="ranking: Entrada, Saída, Cruzado, Análise (cruzado completo + planilha .xlsx) "
                         "ou Diagnóstico de cenários")
    ap.add_argument("--metades", action="store_true",
                    help="modo D: mostra o que se repete nas duas metades do período")
    ap.add_argument("--saida-dir", help="pasta da planilha do modo A (padrão: D:\\DAYTRADE\\ANALISES)")
    ap.add_argument("--periodo", help="'tudo', '01/03/2026 31/03/2026', 'ultimos 40' ou 'mes 2026-03'")
    ap.add_argument("--cenario", help="apura o campeonato só neste cenário (cenario.py): "
                    + ", ".join(cen_mod.NOMES))
    ap.add_argument("--simular", choices=["reamostragem", "regimes"],
                    help="usa dados FICTÍCIOS em vez do histórico real")
    ap.add_argument("--dias-sim", type=int, default=120, help="pregões simulados (padrão 120)")
    ap.add_argument("--semente", type=int, help="semente do simulador (padrão: aleatória, impressa na tela)")
    ap.add_argument("--escala-vol", type=float, default=1.0, help="multiplica a volatilidade simulada")
    ap.add_argument("--sem-espelhar", action="store_true", help="reamostragem sem espelhar pregões")
    return ap.parse_args(argv)


if __name__ == "__main__":
    args = _ler_argumentos()
    if args.simular:
        simulacao = {"modo": args.simular, "dias": args.dias_sim, "semente": args.semente,
                     "escala_vol": args.escala_vol, "espelhar": not args.sem_espelhar}
    elif args.periodo is not None or args.modo:
        simulacao = False       # pediu período/modo na linha de comando: não pergunta a fonte
    else:
        simulacao = None
    cenario = args.cenario
    if cenario is None and (args.modo or args.periodo is not None or args.simular):
        cenario = ""        # rodou por linha de comando: não pergunta, vale todos os cenários
    executar(args.csv, modo=args.modo, periodo=args.periodo, simulacao=simulacao,
             cenario=cenario, saida_dir=args.saida_dir,
             metades=True if args.metades else (False if args.modo else None))
