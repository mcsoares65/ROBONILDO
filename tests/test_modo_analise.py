"""Ranking só por acumulado (sem portões) e planilha do modo A de classificacao.py."""
import contextlib
import io
import tempfile
import sys
import unittest
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import classificacao as cl


def _trade(dia, hora, resultado, lado="COMPRA"):
    dt = datetime(dia.year, dia.month, dia.day, hora, 0)
    return {"horario_rotulo": dt, "horario_execucao": dt, "saida_dt": dt.replace(minute=15),
            "resultado_reais": resultado, "lado": lado, "motivo": "alvo", "intrabar_ambiguo": False,
            "pontos": 0.0}


def _par(entrada, saida, trades, titular_e=False, titular_s=False):
    dias = [date(2026, 9, 1), date(2026, 9, 2), date(2026, 10, 1)]
    r = cl._montar_resultado_par(entrada, saida, trades, dias, entrada_titular=titular_e,
                                 saida_titular=titular_s)
    r["_por_dia"] = cl._resultado_por_dia(trades)
    return r


class RankingSemPortoesTest(unittest.TestCase):
    def test_resumo_nao_tem_pontos_nem_portao_estatistico(self):
        resumo = cl._resumo([_trade(date(2026, 9, 1), 10, 50.0)], 1)
        for chave in ("pontos_dia", "dias_positivos", "sd_operacao", "erro_padrao_total"):
            self.assertNotIn(chave, resumo)
        self.assertFalse(hasattr(cl, "_marcar_elegiveis"))
        self.assertFalse(hasattr(cl, "_rodape_pontos"))

    def test_tabela_sem_coluna_pts_e_com_rodape_do_acumulado(self):
        r = _par("e1", "s1", [_trade(date(2026, 9, 1), 10, 80.0)], True, True)
        r["estrategia"] = "e1"
        r["pos"] = 1
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            cl._imprimir_ranking_simples("T", [r])
        texto = buf.getvalue()
        self.assertNotIn(" pts ", texto)
        self.assertNotIn("Portão", texto)
        self.assertIn("Maior acumulado vence", texto)


class PlanilhaAnaliseTest(unittest.TestCase):
    def setUp(self):
        d1, d2, d3 = date(2026, 9, 1), date(2026, 9, 2), date(2026, 10, 1)
        self.dias = [d1, d2, d3]
        tit = _par("ent_tit", "sai_tit", [_trade(d1, 10, 100.0), _trade(d2, 10, -30.0)], True, True)
        e2 = _par("ent_b", "sai_tit", [_trade(d1, 11, 40.0), _trade(d3, 11, 60.0)], False, True)
        s2 = _par("ent_tit", "sai_b", [_trade(d2, 11, 70.0)], True, False)
        cr = _par("ent_b", "sai_b", [_trade(d3, 12, 25.5)], False, False)
        self.pares = [tit, e2, s2, cr]
        self.rank_e = [dict(r, estrategia=r["entrada"], titular=r["titular_entrada"])
                       for r in self.pares if r["titular_saida"]]
        self.rank_s = [dict(r, estrategia=r["saida"], titular=r["titular_saida"])
                       for r in self.pares if r["titular_entrada"]]
        self.rank_c = [dict(r) for r in self.pares]
        self.meta = {"entrada_titular": "ent_tit", "saida_titular": "sai_tit", "fonte": "histórico completo",
                     "arquivo": "x.csv", "historico_inicio": datetime(2026, 9, 1), "historico_fim": datetime(2026, 10, 1),
                     "n_candles": 100}

    def test_agrupamento_e_resultado_por_dia(self):
        por_dia = cl._resultado_por_dia([_trade(date(2026, 9, 1), 10, 10.0), _trade(date(2026, 9, 1), 11, -4.0)])
        self.assertEqual(por_dia, {date(2026, 9, 1): 6.0})
        mensal = cl._agrupar_periodo({date(2026, 9, 1): 5.0, date(2026, 9, 9): 2.0, date(2026, 10, 1): 1.0},
                                     lambda d: (d.year, d.month))
        self.assertEqual(mensal, {(2026, 9): 7.0, (2026, 10): 1.0})

    def test_colunas_nao_duplicam_pares_com_titular_no_bloco_cruzadas(self):
        cols = cl._colunas_apuracao(self.rank_e, self.rank_s, self.rank_c)
        blocos = [b for b, _, _ in cols]
        self.assertEqual(blocos.count("Entradas"), 2)
        self.assertEqual(blocos.count("Saídas"), 2)
        self.assertEqual([r for b, r, _ in cols if b == "Cruzadas"], ["ent_b × sai_b"])

    def test_planilha_tem_as_abas_e_totais_batem(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as pasta:
            caminho = cl.gravar_planilha_analise(
                Path(pasta) / "a.xlsx", rank_entrada=self.rank_e, rank_saida=self.rank_s,
                rank_cruzado=self.rank_c, dias=self.dias,
                trades_titular=[_trade(self.dias[0], 10, 100.0)], meta=self.meta)
            wb = openpyxl.load_workbook(caminho)
        self.assertEqual(wb.sheetnames[:6], ["Entradas", "Saídas", "Cruzadas", "Apuração diária",
                                             "Apuração mensal", "Apuração anual"])
        diaria = wb["Apuração diária"]
        self.assertEqual(diaria.cell(4, 1).value, "Data")
        self.assertEqual(diaria.max_row, 5 + 3)                       # 3 dias + linha Total
        cabecalhos = [diaria.cell(4, c).value for c in range(3, diaria.max_column + 1)]
        self.assertIn("ent_tit (titular)", cabecalhos)
        col = cabecalhos.index("ent_tit (titular)") + 3
        self.assertEqual(diaria.cell(5, col).value, 100.0)            # 01/09
        self.assertEqual(diaria.cell(8, col).value, 70.0)             # total = 100 - 30
        mensal, anual = wb["Apuração mensal"], wb["Apuração anual"]
        self.assertEqual([mensal.cell(r, 1).value for r in (5, 6)], ["09/2026", "10/2026"])
        self.assertEqual(mensal.cell(5, col - 1).value, 70.0)
        self.assertEqual(anual.cell(5, 1).value, 2026)
        self.assertEqual(anual.cell(5, col - 1).value, 70.0)
        # o total de cada coluna diária é o mesmo da aba de ranking correspondente
        rank = wb["Entradas"]
        resultados = {rank.cell(r, 2).value: rank.cell(r, 3).value for r in range(5, rank.max_row + 1)}
        self.assertEqual(resultados["ent_tit"], 70.0)


if __name__ == "__main__":
    unittest.main()
